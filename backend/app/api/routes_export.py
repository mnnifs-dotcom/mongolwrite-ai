"""Export converted text as downloadable Word (.docx)."""

from __future__ import annotations

import uuid
import zipfile
from io import BytesIO
from pathlib import Path
from xml.etree import ElementTree as ET

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt
from fastapi import APIRouter, HTTPException
from fastapi.responses import Response
from pydantic import BaseModel, Field

router = APIRouter(prefix="/api/v1/export", tags=["export"])

_MAX_CHARS = 200_000
# Bolorsoft face used by KIMO — embedded in the DOCX so Word matches the site.
_PRIMARY_FONT = "Classical Mongolian Dashitseden"
# Traditional Mongolian: top→bottom within a column, columns left→right.
# Matches CSS writing-mode: vertical-lr. Do NOT use tbRl (CJK right→left columns).
_BICHIG_TEXT_DIRECTION = "tbLrV"

_W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
_R_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
_CT_NS = "http://schemas.openxmlformats.org/package/2006/content-types"
_FONT_REL_TYPE = "http://schemas.openxmlformats.org/officeDocument/2006/relationships/font"
_PKG_REL_NS = "http://schemas.openxmlformats.org/package/2006/relationships"

_ASSETS = Path(__file__).resolve().parents[1] / "assets" / "fonts"
_EMBED_FONTS: tuple[tuple[str, Path], ...] = (
    (_PRIMARY_FONT, _ASSETS / "cmdashitseden.ttf"),
    ("MongolianScript", _ASSETS / "MongolianScript.ttf"),
)

ET.register_namespace("w", _W_NS)
ET.register_namespace("r", _R_NS)


class DocxExportRequest(BaseModel):
    text: str = Field(default="", max_length=_MAX_CHARS)
    filename: str = Field(default="mongol-bichig.docx", max_length=120)
    script: str = Field(default="bichig", max_length=20)


def _safe_filename(name: str) -> str:
    cleaned = "".join(ch for ch in name if ch.isalnum() or ch in "._- ").strip() or "mongol-bichig"
    if not cleaned.lower().endswith(".docx"):
        cleaned = f"{cleaned}.docx"
    return cleaned[:120]


def _set_run_font(run, *, size_pt: float) -> None:
    """Mark the run as traditional Mongolian with Dashitseden."""
    run.font.name = _PRIMARY_FONT
    run.font.size = Pt(size_pt)
    r_pr = run._element.get_or_add_rPr()
    r_fonts = r_pr.get_or_add_rFonts()
    # ascii/hAnsi/eastAsia/cs all point at Dashitseden so Word does not
    # substitute a Latin face that breaks Ali Gali shaping.
    for attr in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        r_fonts.set(qn(attr), _PRIMARY_FONT)
    r_fonts.set(qn("w:hint"), "eastAsia")
    # Complex-script size (Word uses szCs for mn-Mong runs).
    sz_cs = r_pr.find(qn("w:szCs"))
    if sz_cs is None:
        sz_cs = OxmlElement("w:szCs")
        r_pr.append(sz_cs)
    sz_cs.set(qn("w:val"), str(int(size_pt * 2)))
    lang = OxmlElement("w:lang")
    lang.set(qn("w:val"), "mn-Mong")
    lang.set(qn("w:eastAsia"), "mn-Mong")
    lang.set(qn("w:bidi"), "mn-Mong")
    r_pr.append(lang)


def _set_section_mongolian_vertical(document: Document) -> None:
    """Top-to-bottom columns progressing left-to-right — Mongolian layout in Word."""
    section = document.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.left_margin = Cm(2.0)
    section.right_margin = Cm(2.0)
    section.top_margin = Cm(2.0)
    section.bottom_margin = Cm(2.0)
    sect_pr = section._sectPr
    for child in list(sect_pr):
        if child.tag == qn("w:textDirection"):
            sect_pr.remove(child)
    text_direction = OxmlElement("w:textDirection")
    text_direction.set(qn("w:val"), _BICHIG_TEXT_DIRECTION)
    sect_pr.append(text_direction)


def _obfuscate_font(data: bytes, font_key: str) -> bytes:
    """ECMA-376 §2.8.1 — XOR first 32 bytes with reversed GUID key."""
    hex_str = "".join(ch for ch in font_key if ch in "0123456789abcdefABCDEF")
    if len(hex_str) != 32:
        raise ValueError(f"Invalid fontKey GUID: {font_key}")
    key = bytearray(16)
    for index in range(16):
        key[15 - index] = int(hex_str[index * 2 : index * 2 + 2], 16)
    out = bytearray(data)
    for index in range(min(32, len(out))):
        out[index] ^= key[index % 16]
    return bytes(out)


def _qname(ns: str, tag: str) -> str:
    return f"{{{ns}}}{tag}"


def _ensure_embed_settings(settings_xml: bytes) -> bytes:
    root = ET.fromstring(settings_xml)
    for tag in ("embedTrueTypeFonts", "saveSubsetFonts"):
        existing = root.find(_qname(_W_NS, tag))
        if existing is None:
            node = ET.SubElement(root, _qname(_W_NS, tag))
            if tag == "embedTrueTypeFonts":
                node.set(_qname(_W_NS, "val"), "true")
            else:
                # Keep full face — Ali Gali glyphs must not be subsetted away.
                node.set(_qname(_W_NS, "val"), "false")
        elif tag == "saveSubsetFonts":
            existing.set(_qname(_W_NS, "val"), "false")
        else:
            existing.set(_qname(_W_NS, "val"), "true")
    return ET.tostring(root, encoding="utf-8", xml_declaration=True)


def _ensure_font_content_types(types_xml: bytes) -> bytes:
    root = ET.fromstring(types_xml)
    defaults = {
        (node.get("Extension") or "").lower(): node
        for node in root.findall(_qname(_CT_NS, "Default"))
    }
    if "odttf" not in defaults:
        node = ET.SubElement(root, _qname(_CT_NS, "Default"))
        node.set("Extension", "odttf")
        node.set(
            "ContentType",
            "application/vnd.openxmlformats-officedocument.obfuscatedFont",
        )
    return ET.tostring(root, encoding="utf-8", xml_declaration=True)


def _build_font_table_and_rels(
    font_table_xml: bytes,
) -> tuple[bytes, bytes, list[tuple[str, bytes]]]:
    """Register embedded Dashitseden (+ MongolianScript) in fontTable + package parts."""
    root = ET.fromstring(font_table_xml)
    keep_names = {_PRIMARY_FONT, "MongolianScript"}
    for font in list(root.findall(_qname(_W_NS, "font"))):
        if font.get(_qname(_W_NS, "name")) in keep_names:
            root.remove(font)

    rel_root = ET.Element(_qname(_PKG_REL_NS, "Relationships"))
    parts: list[tuple[str, bytes]] = []

    for index, (family, path) in enumerate(_EMBED_FONTS, start=1):
        if not path.is_file():
            continue
        font_key = "{" + str(uuid.uuid4()).upper() + "}"
        rel_id = f"rIdEmbed{index}"
        part_name = f"fonts/font{index}.odttf"
        obfuscated = _obfuscate_font(path.read_bytes(), font_key)
        parts.append((f"word/{part_name}", obfuscated))

        font_el = ET.SubElement(root, _qname(_W_NS, "font"))
        font_el.set(_qname(_W_NS, "name"), family)
        charset = ET.SubElement(font_el, _qname(_W_NS, "charset"))
        charset.set(_qname(_W_NS, "val"), "00")
        family_el = ET.SubElement(font_el, _qname(_W_NS, "family"))
        family_el.set(_qname(_W_NS, "val"), "auto")
        pitch = ET.SubElement(font_el, _qname(_W_NS, "pitch"))
        pitch.set(_qname(_W_NS, "val"), "variable")
        embed = ET.SubElement(font_el, _qname(_W_NS, "embedRegular"))
        embed.set(_qname(_R_NS, "id"), rel_id)
        embed.set(_qname(_W_NS, "fontKey"), font_key)

        rel = ET.SubElement(rel_root, _qname(_PKG_REL_NS, "Relationship"))
        rel.set("Id", rel_id)
        rel.set("Type", _FONT_REL_TYPE)
        rel.set("Target", part_name)

    if not parts:
        raise FileNotFoundError("Mongolian font assets missing for Word embed")

    return (
        ET.tostring(root, encoding="utf-8", xml_declaration=True),
        ET.tostring(rel_root, encoding="utf-8", xml_declaration=True),
        parts,
    )


def _embed_mongolian_fonts(docx_bytes: bytes) -> bytes:
    """Inject obfuscated TTF parts so Word renders like the site without system fonts."""
    out = BytesIO()
    with zipfile.ZipFile(BytesIO(docx_bytes), "r") as src, zipfile.ZipFile(
        out, "w", compression=zipfile.ZIP_DEFLATED
    ) as dst:
        names = set(src.namelist())
        base_font_table = (
            src.read("word/fontTable.xml")
            if "word/fontTable.xml" in names
            else (
                b'<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
                b'<w:fonts xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships"'
                b' xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"/>'
            )
        )
        font_table, font_rels, font_parts = _build_font_table_and_rels(base_font_table)

        for name in src.namelist():
            if name in {
                "word/fontTable.xml",
                "word/_rels/fontTable.xml.rels",
            } or name.startswith("word/fonts/"):
                continue
            data = src.read(name)
            if name == "word/settings.xml":
                data = _ensure_embed_settings(data)
            elif name == "[Content_Types].xml":
                data = _ensure_font_content_types(data)
            dst.writestr(name, data)

        dst.writestr("word/fontTable.xml", font_table)
        dst.writestr("word/_rels/fontTable.xml.rels", font_rels)
        for part_name, payload in font_parts:
            dst.writestr(part_name, payload)
    return out.getvalue()


def build_bichig_docx(text: str) -> bytes:
    document = Document()
    _set_section_mongolian_vertical(document)
    lines = text.replace("\r\n", "\n").replace("\r", "\n").split("\n")
    if not any(line.strip() for line in lines):
        lines = [""]
    for line in lines:
        paragraph = document.add_paragraph()
        # Paragraph-level direction reinforces section tbLrV in older Word builds.
        p_pr = paragraph._p.get_or_add_pPr()
        for child in list(p_pr):
            if child.tag == qn("w:textDirection"):
                p_pr.remove(child)
        text_direction = OxmlElement("w:textDirection")
        text_direction.set(qn("w:val"), _BICHIG_TEXT_DIRECTION)
        p_pr.append(text_direction)
        run = paragraph.add_run(line if line else " ")
        _set_run_font(run, size_pt=18)
    buffer = BytesIO()
    document.save(buffer)
    return _embed_mongolian_fonts(buffer.getvalue())


@router.post("/docx")
def export_docx(body: DocxExportRequest) -> Response:
    text = body.text.strip("\ufeff")
    if not text.strip():
        raise HTTPException(status_code=400, detail="Текст хоосон байна")
    if body.script not in {"bichig", "mongol"}:
        raise HTTPException(status_code=400, detail="Зөвхөн монгол бичиг экспортлоно")
    payload = build_bichig_docx(text)
    filename = _safe_filename(body.filename)
    return Response(
        content=payload,
        media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        headers={
            "Content-Disposition": f'attachment; filename="{filename}"',
            "Cache-Control": "no-store",
        },
    )
