from __future__ import annotations

from io import BytesIO
from typing import Annotated

from fastapi import APIRouter, File, HTTPException, UploadFile
from pydantic import BaseModel

router = APIRouter(prefix="/api/v1/import", tags=["import"])

_MAX_BYTES = 8 * 1024 * 1024


class ImportResponse(BaseModel):
    filename: str
    text: str


@router.post("/file", response_model=ImportResponse)
async def import_file(file: Annotated[UploadFile, File()]) -> ImportResponse:
    name = (file.filename or "document").lower()
    data = await file.read()
    if len(data) > _MAX_BYTES:
        raise HTTPException(status_code=413, detail="Файл хэт том байна (8MB).")
    if name.endswith(".docx"):
        text = _docx_to_text(data)
    elif name.endswith(".txt"):
        text = data.decode("utf-8", errors="replace")
    else:
        raise HTTPException(status_code=400, detail="Зөвхөн .docx эсвэл .txt авна.")
    return ImportResponse(filename=file.filename or name, text=text)


def _docx_to_text(data: bytes) -> str:
    from docx import Document

    document = Document(BytesIO(data))
    parts = [paragraph.text for paragraph in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            parts.append("\t".join(cell.text for cell in row.cells))
    return "\n".join(part for part in parts if part.strip())
